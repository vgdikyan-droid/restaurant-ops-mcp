"""Write the categorized videos to an Excel workbook (opens in Excel, Google Sheets, Numbers)."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet

from tiktok_categorizer.categories import UNCATEGORIZED
from tiktok_categorizer.models import SavedVideo

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="25292E")
_LINK_FONT = Font(color="0563C1", underline="single")
_WRAP = Alignment(wrap_text=True, vertical="top")
_TOP = Alignment(vertical="top")
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")

Row = tuple[SavedVideo, list[str]]


def _sheet_title(name: str, taken: set[str]) -> str:
    base = _BAD_SHEET_CHARS.sub("-", name).strip("' ")[:31] or "Category"
    title, n = base, 2
    while title.casefold() in taken:
        suffix = f" ({n})"
        title, n = base[: 31 - len(suffix)] + suffix, n + 1
    taken.add(title.casefold())
    return title


def _text(value: str) -> str:
    return ILLEGAL_CHARACTERS_RE.sub("", value)


def _set_text(cell: Any, value: str) -> None:
    cell.value = _text(value)
    cell.data_type = "s"  # never let a caption like "=..." become a formula


def _posted(value: str | None) -> date | str:
    if not value:
        return ""
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        return value


def _header(sheet: Worksheet, columns: list[tuple[str, int]]) -> None:
    for col, (title, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.freeze_panes = "A2"


def _video_row(sheet: Worksheet, row: int, values: list[Any], video: SavedVideo) -> None:
    for col, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=col)
        if isinstance(value, str):
            _set_text(cell, value)
            cell.alignment = _WRAP
        else:
            cell.value = value
            cell.alignment = _TOP
            if isinstance(value, date):
                cell.number_format = "yyyy-mm-dd"
    link = sheet.cell(row=row, column=len(values) + 1)
    link.hyperlink = video.url
    _set_text(link, video.url)
    link.font, link.alignment = _LINK_FONT, _TOP


def write_workbook(path: Path, rows: list[Row], category_names: list[str]) -> dict[str, int]:
    """Write Summary, All Videos, and one tab per category. Returns videos per category."""
    order = [*category_names, UNCATEGORIZED]
    members: dict[str, list[tuple[int, SavedVideo, list[str]]]] = {name: [] for name in order}
    for number, (video, labels) in enumerate(rows, start=1):
        for label in labels:
            members.setdefault(label, []).append((number, video, labels))
    counts = {name: len(members[name]) for name in members}

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    taken = {"summary", "all videos"}
    tabs = {name: _sheet_title(name, taken) for name in members if members[name]}

    # All Videos: one row per video, in the order they appear in your Favorites.
    everything = workbook.create_sheet("All Videos")
    _header(everything, [("#", 6), ("Category", 22), ("Also in", 28), ("Caption", 70),
                         ("Creator", 22), ("Posted", 12), ("Link", 55)])
    for number, (video, labels) in enumerate(rows, start=1):
        _video_row(everything, number + 1, [number, labels[0], ", ".join(labels[1:]),
                                             video.caption, video.creator, _posted(video.posted_at)], video)
    everything.auto_filter.ref = everything.dimensions

    for name, title in tabs.items():
        sheet = workbook.create_sheet(title)
        _header(sheet, [("#", 6), ("Caption", 70), ("Creator", 22), ("Posted", 12),
                        ("Also in", 28), ("Link", 55)])
        for row, (number, video, labels) in enumerate(members[name], start=2):
            others = ", ".join(label for label in labels if label != name)
            _video_row(sheet, row, [number, video.caption, video.creator,
                                    _posted(video.posted_at), others], video)
        sheet.auto_filter.ref = sheet.dimensions

    summary["A1"] = "Saved TikTok videos by category"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A2"] = (f"{len(rows)} saved videos, sorted {datetime.now():%Y-%m-%d %H:%M}. "
                     "A video can appear in more than one category. Click a category to open its tab.")
    for col, (title, width) in enumerate([("Category", 30), ("Videos", 10)], start=1):
        cell = summary.cell(row=4, column=col, value=title)
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
        summary.column_dimensions[cell.column_letter].width = width
    for row, name in enumerate(members, start=5):
        cell = summary.cell(row=row, column=1)
        _set_text(cell, name)
        summary.cell(row=row, column=2, value=counts[name])
        if name in tabs:
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{tabs[name]}'!A1")
            cell.font = _LINK_FONT

    workbook.save(path)
    return counts
