from contextlib import contextmanager
import json

from openpyxl import load_workbook

from tiktok_categorizer import cli
from tiktok_categorizer.captions import fill_missing_captions, lookup_caption
from tiktok_categorizer.categories import parse_categories
from tiktok_categorizer.models import SavedVideo, load_video_list, save_video_list

URL = "https://www.tiktok.com/@janestyle/video/7301234567890123456"
OEMBED = {"title": "fit check #ootd", "author_name": "Jane", "author_url": "https://www.tiktok.com/@janestyle"}


def test_oembed_fills_in_caption_and_creator():
    video = SavedVideo(url=URL)
    seen = []

    def http_get(url):
        seen.append(url)
        return 200, json.dumps(OEMBED).encode()

    assert lookup_caption(video, http_get)
    assert (video.caption, video.creator) == ("fit check #ootd", "@janestyle")
    assert seen == ["https://www.tiktok.com/oembed?url=https%3A%2F%2Fwww.tiktok.com%2F%40janestyle%2Fvideo%2F7301234567890123456"]


def test_oembed_retries_rate_limits_but_not_missing_videos():
    replies = [(429, b""), (200, json.dumps(OEMBED).encode())]
    sleeps = []
    assert lookup_caption(SavedVideo(url=URL), lambda _: replies.pop(0), sleep=sleeps.append)
    assert sleeps == [1]

    calls = []
    assert not lookup_caption(SavedVideo(url=URL), lambda u: calls.append(u) or (404, b""), sleep=sleeps.append)
    assert len(calls) == 1


def test_only_videos_with_nothing_known_are_looked_up():
    videos = [SavedVideo(url=URL), SavedVideo(url=URL, creator="@x"), SavedVideo(url=URL, caption="hi")]
    calls = []

    def http_get(url):
        calls.append(url)
        return 200, json.dumps(OEMBED).encode()

    assert fill_missing_captions(videos, http_get) == 1
    assert len(calls) == 1


def _saved(tmp_path):
    path = tmp_path / "list.json"
    save_video_list(path, [
        SavedVideo(url=URL, caption="fit check #ootd", creator="@janestyle", posted_at="2023-11-14"),
        SavedVideo(url=URL.replace("3456", "3457"), caption="donate to the #gofundme", creator="@helper"),
        SavedVideo(url=URL.replace("3456", "3458"), caption="", creator="@quiet"),
    ], "me")
    return path


def test_resorting_a_saved_list_writes_the_spreadsheet(tmp_path, capsys):
    out = tmp_path / "videos.xlsx"
    assert cli.main(["--from-list", str(_saved(tmp_path)), "-o", str(out)]) == 0

    book = load_workbook(out)
    assert {"Fashion", "Fundraisers & Causes", "Uncategorized"} <= set(book.sheetnames)
    assert "Fashion: 1" in capsys.readouterr().err
    assert not (tmp_path / "videos.json").exists()  # the input list is not rewritten


def test_collecting_from_tiktok_saves_the_list_for_later(tmp_path, monkeypatch):
    calls = {}

    @contextmanager
    def fake_browser(profile_dir, **kwargs):
        calls["profile_dir"] = profile_dir
        yield "context", "page"

    def fake_collect(page, username, *, limit, log):
        calls["collect"] = (page, username, limit)
        return load_video_list(_saved(tmp_path))

    monkeypatch.setattr(cli, "open_tiktok_browser", fake_browser)
    monkeypatch.setattr(cli, "wait_for_login", lambda context, page, log: None)
    monkeypatch.setattr(cli, "detect_username", lambda page: "detected_me")
    monkeypatch.setattr(cli, "collect_favorites", fake_collect)

    out = tmp_path / "mine.xlsx"
    assert cli.main(["-o", str(out), "--limit", "5", "--profile-dir", str(tmp_path / "profile")]) == 0
    assert calls["collect"] == ("page", "detected_me", 5)
    assert len(load_video_list(tmp_path / "mine.json")) == 3
    assert out.exists()


def test_ai_mode_without_the_package_or_key_falls_back(tmp_path, monkeypatch, capsys):
    from tiktok_categorizer import ai

    def unavailable(*args, **kwargs):
        raise ai.AIUnavailable("no key")

    monkeypatch.setattr(ai, "AIClassifier", unavailable)
    out = tmp_path / "videos.xlsx"
    assert cli.main(["--from-list", str(_saved(tmp_path)), "-o", str(out), "--ai"]) == 0
    assert "keyword rules instead" in capsys.readouterr().err
    assert out.exists()


def test_default_categories_print_as_a_usable_file(capsys):
    assert cli.main(["--print-default-categories"]) == 0
    assert parse_categories(json.loads(capsys.readouterr().out))


def test_bad_inputs_exit_with_a_message(tmp_path, capsys):
    bad = tmp_path / "cats.json"
    bad.write_text("[]")
    assert cli.main(["--from-list", str(_saved(tmp_path)), "--categories", str(bad)]) == 1
    assert cli.main(["--from-list", str(tmp_path / "missing.json")]) == 1
    assert cli.main(["--limit", "0"]) == 2
    err = capsys.readouterr().err
    assert "Error:" in err and "--limit" in err
