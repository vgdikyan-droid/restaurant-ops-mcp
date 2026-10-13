from datetime import datetime

from openpyxl import load_workbook

from tiktok_categorizer.categories import UNCATEGORIZED
from tiktok_categorizer.models import SavedVideo
from tiktok_categorizer.spreadsheet import write_workbook


def _video(n, caption, posted="2024-05-01"):
    return SavedVideo(url=f"https://www.tiktok.com/@c{n}/video/71000000000000000{n:02d}",
                      caption=caption, creator=f"@c{n}", posted_at=posted)


def test_workbook_has_a_tab_per_category_with_links(tmp_path):
    path = tmp_path / "out.xlsx"
    rows = [
        (_video(1, "fit check #ootd"), ["Fashion"]),
        (_video(2, "thrift haul for charity"), ["Fashion", "Fundraisers & Causes"]),
        (_video(3, "lol", posted=None), [UNCATEGORIZED]),
    ]
    counts = write_workbook(path, rows, ["Fashion", "Fundraisers & Causes", "Cool Edits"])
    assert counts == {"Fashion": 2, "Fundraisers & Causes": 1, "Cool Edits": 0, UNCATEGORIZED: 1}

    book = load_workbook(path)
    # Empty categories get no tab; "&" is fine in a tab name.
    assert book.sheetnames == ["Summary", "All Videos", "Fashion", "Fundraisers & Causes", UNCATEGORIZED]

    fashion = book["Fashion"]
    assert [c.value for c in fashion[1]] == ["#", "Caption", "Creator", "Posted", "Also in", "Link"]
    assert fashion["B3"].value == "thrift haul for charity"
    assert fashion["E3"].value == "Fundraisers & Causes"
    assert fashion["F2"].value == rows[0][0].url
    assert fashion["F2"].hyperlink.target == rows[0][0].url
    assert fashion["D2"].value == datetime(2024, 5, 1)

    everything = book["All Videos"]
    assert [everything.cell(row=r, column=2).value for r in (2, 3, 4)] == ["Fashion", "Fashion", UNCATEGORIZED]
    assert everything["C3"].value == "Fundraisers & Causes"

    summary = book["Summary"]
    assert summary["A5"].value == "Fashion" and summary["B5"].value == 2
    assert summary["A5"].hyperlink.location == "'Fashion'!A1"
    assert summary["A7"].value == "Cool Edits" and summary["A7"].hyperlink is None


def test_captions_cannot_become_formulas_or_break_the_file(tmp_path):
    path = tmp_path / "out.xlsx"
    caption = '=HYPERLINK("http://evil.example","click")\x07'
    write_workbook(path, [(_video(1, caption), ["Fashion"])], ["Fashion"])
    cell = load_workbook(path)["Fashion"]["B2"]
    assert cell.data_type == "s"
    assert cell.value == '=HYPERLINK("http://evil.example","click")'


def test_awkward_category_names_make_valid_unique_tabs(tmp_path):
    path = tmp_path / "out.xlsx"
    names = ["Work/Career: 9-5 [office]", "Summary", "A" * 40, "A" * 40 + "!"]
    rows = [(_video(i, "x"), [name]) for i, name in enumerate(names)]
    write_workbook(path, rows, names)
    sheetnames = load_workbook(path).sheetnames
    assert len(sheetnames) == 2 + len(names)
    assert "Work-Career- 9-5 -office-" in sheetnames
    assert all(len(name) <= 31 for name in sheetnames)
    assert len({name.casefold() for name in sheetnames}) == len(sheetnames)
